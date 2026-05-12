import json
import copy
from io import BytesIO
from pathlib import Path

import openpyxl

from ..config import settings
from ..models.character import CharacterCard

EXCEL_TEMPLATE = settings.ORIGIN_DATA_DIR / "狩魂者空白卡正式版V1.4.xlsx"

CELL_MAPPING = {
    "角色卡": {
        "basic_info.name": (15, 2),
        "basic_info.codename": (15, 10),
        "basic_info.gender": (16, 2),
        "basic_info.age": (16, 10),
        "basic_info.bond": (17, 2),
        "spirit_power.name": (7, 17),
        "spirit_power.description": (9, 17),
        "attributes.physique": (15, 17),
        "attributes.wisdom": (17, 17),
        "attributes.spirit": (19, 17),
        "resources.spirit_awareness": (24, 17),
    },
}


def _get_nested_value(obj: dict, path: str):
    parts = path.split(".")
    current = obj
    for part in parts:
        if isinstance(current, dict):
            current = current.get(part, "")
        else:
            return ""
    return current if current is not None else ""


def export_json(card: CharacterCard) -> dict:
    return card.model_dump()


def export_excel(card: CharacterCard) -> bytes:
    if not EXCEL_TEMPLATE.exists():
        raise FileNotFoundError(f"Excel template not found: {EXCEL_TEMPLATE}")

    wb = openpyxl.load_workbook(str(EXCEL_TEMPLATE))
    card_data = card.model_dump()

    for sheet_name, mappings in CELL_MAPPING.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for field_path, (row, col) in mappings.items():
            value = _get_nested_value(card_data, field_path)
            if value:
                ws.cell(row=row, column=col, value=str(value))

    ws = wb["角色卡"]

    if card.martial_arts:
        base_row = 41
        for i, art in enumerate(card.martial_arts[:5]):
            r = base_row + i * 2
            ws.cell(row=r, column=2, value=art.name)
            ws.cell(row=r, column=8, value=art.cost)
            ws.cell(row=r, column=13, value=art.target)
            ws.cell(row=r, column=20, value=art.effect)

    if card.spells:
        base_row = 41
        for i, spell in enumerate(card.spells[:5]):
            r = base_row + i * 2
            ws.cell(row=r, column=44, value=spell.name)
            ws.cell(row=r, column=50, value=spell.cost)
            ws.cell(row=r, column=55, value=spell.target)
            ws.cell(row=r, column=62, value=spell.spell_type)
            ws.cell(row=r, column=65, value=spell.effect)

    if card.feats:
        for i, feat in enumerate(card.feats[:2]):
            col_offset = 0 if i == 0 else 42
            ws.cell(row=28, column=2 + col_offset, value=feat.name)
            ws.cell(row=28, column=10 + col_offset, value=feat.description)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def import_json(data: dict) -> CharacterCard:
    return CharacterCard(**data)
